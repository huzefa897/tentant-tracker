import json
import tempfile
from datetime import date, datetime
from io import StringIO
from pathlib import Path
from unittest import skipUnless

from django.core.management import CommandError, call_command
from django.test import TestCase

from billing.legacy_import import import_ledger
from billing.legacy_ledger import parse_ledger
from billing.models import Receipt, RentEntry

from .factories import make_property, make_tenancy, make_user
from .test_payments_api import TODAY

try:
    import openpyxl
except ImportError:  # A dev package: production images don't have it.
    openpyxl = None

if openpyxl is not None:
    from billing.sheet_ledger import convert_workbook


def tenant_tab(workbook, title, rows, *, rate=6, flat=202.0, notes_header="Notes"):
    """A tab laid out like the owners' sheet (fictional data): details, then monthly rows."""
    sheet = workbook.create_sheet(title)
    sheet["A1"] = "Rent Tracker"
    sheet["A3"], sheet["C3"] = "Name of Tenant", title
    sheet["A4"], sheet["C4"] = "Rent Amount", 15000
    sheet["D4"], sheet["F4"] = "Flat Number", flat
    sheet["A5"], sheet["C5"] = "Rent Due Day (of month)", "11th "
    sheet["D5"], sheet["F5"] = "Adhaar Card Number", "0000 0000 0000"
    sheet["A6"], sheet["C6"] = "Electricity Per Unit Cost", rate
    sheet.append([])
    sheet.append([])
    sheet.append([])
    # Row 10: the header, in this tab's own order.
    headers = ["Month", "Rent Amount", "Status", "Units Used", "Electricity Charge",
               "Water Bill", "Total Due", "Total Paid", "Paid Date", notes_header,
               "Email Trigger", "Overdue Amount"]
    for column, header in enumerate(headers, start=1):
        sheet.cell(row=10, column=column, value=header)
    for number, values in enumerate(rows, start=11):
        for column, value in enumerate(values, start=1):
            sheet.cell(row=number, column=column, value=value)
    return sheet


def row(month, units, paid=None, paid_date=None, *, rent=15000, water=None, notes=None,
        total_due=None, charge=None, rate=6, overdue=0):
    # Total Due and Overdue Amount are never imported, only compared. By default Total Due is
    # the month's own charges, with nothing carried forward from last month.
    charge = round(units * rate, 2) if charge is None and units is not None else charge
    if total_due is None:
        total_due = rent + (charge or 0) + (water or 0)
    return [month, rent, "Paid", units, charge, water, total_due, paid, paid_date, notes,
            False, overdue]


def workbook_with(*tabs):
    workbook = openpyxl.Workbook()
    workbook.active.title = "README"
    workbook.create_sheet("Template")
    for title, rows, options in tabs:
        tenant_tab(workbook, title, rows, **options)
    workbook.create_sheet("Email Log")
    return workbook


def convert(*tabs):
    return convert_workbook(workbook_with(*tabs), today=TODAY)


@skipUnless(openpyxl, "needs openpyxl (pipenv install --dev)")
class ConvertWorkbookTests(TestCase):
    def test_converts_each_months_own_charges_and_payment(self):
        conversion = convert(
            (
                "Tenant A",
                [
                    row(datetime(2026, 8, 1), 75, 15750, datetime(2026, 8, 16), water=300),
                    # Text date, a placeholder note, part paid.
                    row(datetime(2026, 9, 1), 80.5, 5000, "20/09/2026", notes="None"),
                ],
                {},
            )
        )

        self.assertEqual(conversion.problems, [])
        self.assertEqual(conversion.skipped_tabs, ["README", "Template", "Email Log"])
        tenant = conversion.ledger["tenants"][0]
        self.assertEqual(tenant["sourceId"], "Tenant A")
        august, september = tenant["bills"]
        self.assertEqual(
            august,
            {
                "sourceId": "Tenant A/2026-08",
                "billingMonth": "2026-08",
                "rentAmount": "15000.00",
                "electricityUnits": "75.000",
                "electricityRate": "6.00",
                "waterAmount": "300.00",
                "notes": "",
                "payments": [
                    {
                        "sourceId": "Tenant A/2026-08/payment",
                        "amount": "15750.00",
                        "paidDate": "2026-08-16",
                        "notes": "",
                        "receipt": {"reference": "RCP-202-202608", "issuedDate": "2026-08-16"},
                    }
                ],
            },
        )
        self.assertEqual(september["electricityUnits"], "80.500")
        self.assertEqual(september["notes"], "")
        self.assertEqual(september["payments"][0]["paidDate"], "2026-09-20")

    def test_never_reads_aadhaar_numbers(self):
        conversion = convert(("Tenant A", [row(datetime(2026, 8, 1), 75)], {}))

        self.assertNotIn("0000 0000 0000", json.dumps(conversion.ledger))

    def test_leaves_out_future_and_unfilled_months(self):
        conversion = convert(
            (
                "Tenant A",
                [
                    row(datetime(2026, 9, 1), 75),
                    row(datetime(2026, 10, 1), None),
                    row(datetime(2026, 11, 1), 10),
                ],
                {},
            )
        )

        self.assertEqual(len(conversion.ledger["tenants"][0]["bills"]), 1)
        self.assertEqual(len(conversion.warnings), 2)

    def test_refuses_a_payment_covering_other_months(self):
        # ₹31,500 paid against ₹15,750: the sheet's carry-forward, which the app doesn't do.
        conversion = convert(
            ("Tenant A", [row(datetime(2026, 9, 1), 75, 31500, datetime(2026, 9, 5), water=300)],
             {})
        )

        self.assertIn("move the part that paid other months", conversion.problems[0])
        self.assertIn("'Tenant A' row 11", conversion.problems[0])

    def test_refuses_a_payment_without_a_date(self):
        conversion = convert(("Tenant A", [row(datetime(2026, 9, 1), 75, 1000, "soon")], {}))

        self.assertIn("Paid Date 'soon' isn't a date", conversion.problems[0])

    def test_warns_about_dates_that_look_swapped(self):
        # An August bill "paid" on 8 January.
        conversion = convert(
            ("Tenant A", [row(datetime(2026, 8, 1), 75, 1000, datetime(2026, 1, 8))], {})
        )

        self.assertIn("check the day and month aren't swapped", conversion.warnings[0])

    def test_keeps_a_typed_over_electricity_charge_when_it_is_a_whole_rate(self):
        conversion = convert(("Tenant A", [row(datetime(2026, 9, 1), 100, charge=700)], {}))

        self.assertEqual(conversion.ledger["tenants"][0]["bills"][0]["electricityRate"], "7.00")
        self.assertIn("electricity charged at 7.00/unit", conversion.warnings[0])

    def test_follows_the_sheets_carry_forward_without_warning(self):
        # August left ₹5,750 owing; September's Total Due adds it once, as it should.
        conversion = convert(
            (
                "Tenant A",
                [
                    row(datetime(2026, 8, 1), 75, 10000, datetime(2026, 8, 16), water=300,
                        overdue=5750),
                    row(datetime(2026, 9, 1), 80, total_due=15480 + 5750, overdue=21230),
                ],
                {},
            )
        )

        self.assertEqual(conversion.problems, [])
        self.assertEqual(conversion.warnings, [])

    def test_warns_once_per_tab_where_total_due_isnt_the_months_charges(self):
        conversion = convert(
            (
                "Tenant A",
                [
                    row(datetime(2026, 8, 1), 75, 10000, datetime(2026, 8, 16), water=300,
                        overdue=5750),
                    # Last month's balance added twice, as on two of the owners' tabs.
                    row(datetime(2026, 9, 1), 80, total_due=15480 + 2 * 5750, overdue=26980),
                    # Water left out of Total Due.
                    row(datetime(2026, 10, 1), 70, water=200, total_due=15420 + 26980),
                ],
                {},
            )
        )

        self.assertEqual(conversion.problems, [])
        self.assertEqual(len(conversion.warnings), 1)
        warning = conversion.warnings[0]
        self.assertIn("2 month(s)", warning)
        self.assertIn("(rows 12, 13)", warning)
        self.assertIn("September 2026 shows 21230.00, the ledger bills 15480.00", warning)
        # Only compared: the ledger keeps the month's own charges.
        october = conversion.ledger["tenants"][0]["bills"][2]
        self.assertEqual(october["waterAmount"], "200.00")

    def test_skips_tabs_without_a_month_header(self):
        workbook = workbook_with()
        workbook.create_sheet("Notes")["A1"] = "Ideas"

        conversion = convert_workbook(workbook, today=TODAY)

        self.assertIn("Notes", conversion.skipped_tabs)
        self.assertEqual(conversion.ledger["tenants"], [])

    def test_its_ledger_imports(self):
        conversion = convert(
            (
                "Tenant A",
                [
                    row(datetime(2026, 8, 1), 75, 15750, datetime(2026, 8, 16), water=300),
                    row(datetime(2026, 9, 1), 80, 5000, "20/09/2026"),
                ],
                {},
            )
        )
        user = make_user()
        tenancy = make_tenancy(make_property(member=user), start_date=date(2026, 1, 1))

        report = import_ledger(
            parse_ledger(conversion.ledger, today=TODAY),
            {"Tenant A": tenancy.pk},
            imported_by=user,
            dry_run=False,
        )

        self.assertTrue(report.saved)
        self.assertEqual(RentEntry.objects.count(), 2)
        self.assertEqual(
            sorted(Receipt.objects.values_list("reference", flat=True)),
            ["RCP-202-202608", "RCP-202-202609"],
        )
        self.assertIn(("total paid", "20750.00", "20750.00"), report.reconciliation)


@skipUnless(openpyxl, "needs openpyxl (pipenv install --dev)")
class SheetToLedgerCommandTests(TestCase):
    def setUp(self):
        folder = tempfile.TemporaryDirectory()
        self.addCleanup(folder.cleanup)
        self.folder = Path(folder.name)
        self.workbook = self.folder / "TenantTracker.xlsx"

    def save(self, rows):
        workbook_with(("Tenant A", rows, {})).save(self.workbook)

    def call(self, *extra):
        out, err = StringIO(), StringIO()
        call_command(
            "sheet_to_ledger",
            str(self.workbook),
            "--output",
            str(self.folder / "ledger.json"),
            *extra,
            stdout=out,
            stderr=err,
        )
        return out.getvalue()

    def test_writes_the_ledger_and_a_mapping_template(self):
        self.save([row(datetime(2026, 9, 1), 75, 1000, datetime(2026, 9, 5))])

        output = self.call("--mapping-template", str(self.folder / "map.json"))

        ledger = json.loads((self.folder / "ledger.json").read_text())
        self.assertEqual(ledger["tenants"][0]["bills"][0]["billingMonth"], "2026-09")
        self.assertEqual(json.loads((self.folder / "map.json").read_text()), {"Tenant A": ""})
        self.assertIn("never commit it", output)

    def test_writes_nothing_when_there_are_problems(self):
        self.save([row(datetime(2026, 9, 1), 75, 1000, "soon")])

        with self.assertRaisesMessage(CommandError, "Nothing was written"):
            self.call()

        self.assertFalse((self.folder / "ledger.json").exists())

    def test_never_overwrites_without_force(self):
        self.save([row(datetime(2026, 9, 1), 75)])
        (self.folder / "ledger.json").write_text("{}")

        with self.assertRaisesMessage(CommandError, "pass --force"):
            self.call()
